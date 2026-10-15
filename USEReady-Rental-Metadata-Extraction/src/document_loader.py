from pathlib import Path
from docx import Document
from PIL import Image, ImageOps, ImageEnhance, ImageFilter
import pytesseract
from .config import TESSERACT_CMD

if TESSERACT_CMD:
    pytesseract.pytesseract.tesseract_cmd = TESSERACT_CMD


def load_docx(path: Path) -> str:
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        if p.text.strip():
            parts.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            vals = [cell.text.strip() for cell in row.cells]
            vals = [v for v in vals if v]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def load_image(path: Path) -> str:
    image = Image.open(path).convert("L")
    # Mild OCR preprocessing; no document-specific extraction rules are used.
    scale = 2
    image = image.resize((image.width * scale, image.height * scale))
    image = ImageOps.autocontrast(image)
    image = ImageEnhance.Sharpness(image).enhance(1.5)
    image = image.filter(ImageFilter.SHARPEN)
    return pytesseract.image_to_string(image, config="--psm 6")


def load_document(path: str | Path) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text = load_docx(path)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff"}:
        text = load_image(path)
    else:
        raise ValueError(f"Unsupported document type: {path.suffix}")
    return " ".join(text.split())
